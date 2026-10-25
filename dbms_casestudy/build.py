# -*- coding: utf-8 -*-
import os, sys, glob
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import *

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SHOTS = sys.argv[1] if len(sys.argv) > 1 else None
OUT   = sys.argv[2] if len(sys.argv) > 2 else "report.docx"

doc = Document()

# ---- page + base styles -------------------------------------------------
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.9)
    s.left_margin = s.right_margin = Inches(1.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

def shade(el, hexcolor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    el.append(sh)

def border(p, color="BFBFBF", sz="6"):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "6"); e.set(qn("w:color"), color)
        bd.append(e)
    pPr.append(bd)

def para(text="", size=11, bold=False, italic=False, align=None,
         before=0, after=8, color=None, font=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size)
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return p

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string("1F3864")
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8")
    e.set(qn("w:space"), "4"); e.set(qn("w:color"), "1F3864")
    bd.append(e); pPr.append(bd)
    return p

def h2(text):
    return para(text, size=12.5, bold=True, before=14, after=6, color="2E5496")

def h3(text):
    return para(text, size=11.5, bold=True, before=12, after=4, color="333333")

def code_block(code):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.0
    lines = code.split("\n")
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string("1A1A1A")
        if i < len(lines) - 1:
            r.add_break()
    shade(p._p.get_or_add_pPr(), "F2F2F2")
    border(p, "D0D0D0")
    return p

def mk_table(rows, widths=None, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cp = cells[ci].paragraphs[0]
            cp.paragraph_format.space_before = Pt(3)
            cp.paragraph_format.space_after = Pt(3)
            r = cp.add_run(val)
            r.font.size = Pt(9.5)
            if header and ri == 0:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string("FFFFFF")
                shade(cells[ci]._tc.get_or_add_tcPr(), "2E5496")
            elif ci == 0:
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

FIGN = [0]
def figure(slot, caption):
    FIGN[0] += 1
    img = None
    if SHOTS and os.path.isdir(SHOTS):
        hits = sorted(glob.glob(os.path.join(SHOTS, slot + "*")))
        hits = [h for h in hits if h.lower().endswith((".png", ".jpg", ".jpeg", ".gif"))]
        if hits: img = hits[0]
    if img:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(img, width=Inches(6.0))
    else:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        for i, txt in enumerate([f"[ screenshot {slot} ]", caption,
                                 f"paste the mongosh output here, or save it as {slot}.png"]):
            r = p.add_run(txt)
            r.font.size = Pt(9 if i else 10)
            r.italic = True
            r.font.color.rgb = RGBColor.from_string("888888")
            if i < 2: r.add_break()
        shade(p._p.get_or_add_pPr(), "FAFAFA")
        border(p, "C0C0C0", sz="4")
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(14)
    r = c.add_run(f"Fig. {FIGN[0]} — {caption}")
    r.italic = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string("555555")

def blocks(items):
    for kind, val in items:
        if kind == "p": para(val, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "h": h2(val)
        elif kind == "h3": h3(val)
        elif kind == "table2": mk_table(val, widths=[1.3, 5.2])

# =========================================================================
# Title page
# =========================================================================
para("", after=40)
para(STUDENT["course"].upper(), size=18, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=4, color="1F3864")
para(STUDENT["sem"], size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Case Study using MongoDB — " + STUDENT["slot"], size=13,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=46)

para("NEWS PORTAL", size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
     after=6, color="2E5496")
para("A document-oriented database for articles, authors and readers",
     size=11.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=56)

t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in [("Name", STUDENT["name"]), ("Registration number", STUDENT["reg"]),
             ("Case study topic", STUDENT["topic"]),
             ("Database", STUDENT["db"]),
             ("Collections", "categories, authors, articles, readers, comments")]:
    cells = t.add_row().cells
    cells[0].width = Inches(1.9); cells[1].width = Inches(4.2)
    r0 = cells[0].paragraphs[0].add_run(k); r0.bold = True; r0.font.size = Pt(10.5)
    shade(cells[0]._tc.get_or_add_tcPr(), "EDF2F9")
    r1 = cells[1].paragraphs[0].add_run(v); r1.font.size = Pt(10.5)

doc.add_page_break()

# =========================================================================
h1("1. Case study understanding")
blocks(TASK1)

h1("2. Collection identification")
blocks(TASK2)

h1("3. Document design")
blocks(TASK3_INTRO)
for name, rows in TASK3_TABLES:
    h2(name)
    mk_table([["Field", "Type", "Description"]] + rows, widths=[1.5, 1.6, 3.4])
blocks(TASK3_OUTRO)

doc.add_page_break()
h1("4. Database creation")
blocks(TASK4_INTRO)

def emit(slot_ids):
    for o in OPS:
        if o["slot"] not in slot_ids: continue
        h2(o["heading"])
        code_block(o["code"])
        para(o["expl"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        figure(o["slot"], o["caption"])

emit({"S01"})

h1("5. Create operations")
blocks(TASK5_INTRO)
emit({f"S0{i}" for i in range(2, 9)})

h1("6. Read operations")
blocks(TASK6_INTRO)
emit({"S09"} | {f"S{i}" for i in range(10, 21)})

h1("7. Update operations")
blocks(TASK7_INTRO)
emit({f"S{i}" for i in range(21, 27)})

h1("8. Delete operations")
blocks(TASK8_INTRO)
emit({f"S{i}" for i in range(27, 30)})

h1("Summary of operations performed")
mk_table([
    ["Task", "Required", "Performed"],
    ["Collections", "5 minimum", "5 — categories, authors, articles, readers, comments"],
    ["Create", "8-10 documents per major collection", "8 categories, 8 authors, 10 articles, 8 readers, 10 comments"],
    ["Read", "8 minimum", "12 — R1 to R12, including two aggregation pipelines"],
    ["Update", "4 minimum", "6 — U1 to U6, using $inc, $set, $currentDate, $addToSet and $push"],
    ["Delete", "2 minimum", "3 — one deleteOne and two deleteMany"],
], widths=[1.3, 2.0, 3.2])

doc.save(OUT)
print("saved", OUT, "| figures:", FIGN[0])
